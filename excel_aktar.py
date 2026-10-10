"""
Plan Excel'ini okuyup veritabanına yazar.

Beklenen düzen (örnek):
    3. satır (başlığın bir üstü) :            2026 AĞUSTOS            ...  (ay başlığı, birleştirilmiş olabilir)
    4. satır (başlık)            : Firma | REFERANS | ... | Backlog | Güncel Backlog | Saat | W34 | Saat | W35 | Saat | Stok Durumu
    5. satırdan itibaren         : ürünler

- Başlık satırı "REFERANS" yazan satırdır (ilk 20 satırda aranır).
- Ay başlığı başlık satırının hemen üstündedir; içinde yıl yazıyorsa ("2026 AĞUSTOS") okunur.
  Yıl yoksa kullanıcıya sorulur (tahmin edilmez).
- Bir dosyada birden fazla ay olabilir; her ay ayrı bir yükleme (işlem kaydı) olur.
- Backlog / Stok Durumu dosyadaki İLK aya yazılır.
- "Saat", "Toplam" vb. sütunlar okunmaz; saat, uygulamada operasyon saniyelerinden hesaplanır.
"""
import os
import re
from dataclasses import dataclass, field
from datetime import datetime
from typing import Dict, List, Optional

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter

from capacity_engine import normalize_code
from veritabani import connect, create_empty_db, hafta_yili

AYLAR = ["OCAK", "ŞUBAT", "MART", "NİSAN", "MAYIS", "HAZİRAN",
         "TEMMUZ", "AĞUSTOS", "EYLÜL", "EKİM", "KASIM", "ARALIK"]


def fold(text) -> str:
    """Karşılaştırma için: büyük harf + Türkçe karakterleri sadeleştir (Ağustos -> AGUSTOS)."""
    s = str(text).strip()
    s = s.replace("i", "İ").replace("ı", "I").upper()
    for a, b in zip("ÇĞİÖŞÜÂÎÛ", "CGIOSUAIU"):
        s = s.replace(a, b)
    return re.sub(r"\s+", " ", s)


AY_FOLD = {fold(a): a for a in AYLAR}


def parse_month_label(text):
    """'2026 AĞUSTOS' -> ('AĞUSTOS', 8, 2026); 'Eylül' -> ('EYLÜL', 9, None); ay yoksa None."""
    if text is None or (isinstance(text, float) and pd.isna(text)):
        return None
    if isinstance(text, (datetime, pd.Timestamp)):
        return AYLAR[text.month - 1], text.month, text.year
    f = fold(text)
    ay = next((AY_FOLD[k] for k in AY_FOLD if re.search(rf"\b{k}\b", f)), None)
    if not ay:
        return None
    m = re.search(r"\b(20\d\d)\b", f)
    return ay, AYLAR.index(ay) + 1, int(m.group(1)) if m else None


@dataclass
class MonthBlock:
    label: str                     # Excel'de yazan
    ay: str
    ay_no: int
    yil: Optional[int]             # None -> kullanıcı girecek
    weeks: Dict[int, int] = field(default_factory=dict)   # sütun index -> hafta no


@dataclass
class ParsedPlan:
    path: str
    sheet: str
    blocks: List[MonthBlock]
    orders: pd.DataFrame           # urun_kodu, firma, blok, hafta, adet
    backlog: pd.DataFrame          # urun_kodu, firma, backlog_adet, stok_adet
    n_products: int
    warnings: List[str]


try:   # eski .xls dosyaları için (exe'ye dahil edilsin diye burada açıkça içe aktarılır)
    import xlrd  # noqa: F401
    HAS_XLRD = True
except ImportError:
    HAS_XLRD = False


def _check_format(path: str):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xls" and not HAS_XLRD:
        raise ValueError("Bu bir eski Excel (.xls) dosyası ve okumak için gereken 'xlrd' paketi kurulu değil.\n"
                         "Çözüm: build_exe.bat'ı yeniden çalıştırın (xlrd'yi kurar) ya da dosyayı Excel'de "
                         "'Farklı Kaydet > Excel Çalışma Kitabı (.xlsx)' ile kaydedin.")
    if ext not in (".xlsx", ".xlsm", ".xls"):
        raise ValueError(f"Desteklenmeyen dosya türü: {ext or '(uzantı yok)'}. .xlsx, .xlsm veya .xls seçin.")


def read_sheet(path: str, sheet: str) -> pd.DataFrame:
    """Sayfayı başlıksız okur (.xlsx/.xlsm: openpyxl, .xls: xlrd)."""
    _check_format(path)
    return pd.read_excel(path, sheet_name=sheet, header=None)


def sheet_names(path: str) -> List[str]:
    _check_format(path)
    return pd.ExcelFile(path).sheet_names


def parse_plan(path: str, sheet: str) -> ParsedPlan:
    raw = read_sheet(path, sheet)
    warnings = []

    # --- başlık satırı ---
    hdr = None
    for r in range(min(20, len(raw))):
        if any(fold(v) == "REFERANS" for v in raw.iloc[r] if pd.notna(v)):
            hdr = r
            break
    if hdr is None:
        raise ValueError("Başlık satırı bulunamadı: ilk 20 satırda 'REFERANS' yazan hücre yok.")
    if hdr == 0:
        raise ValueError("Ay başlığı bulunamadı: 'REFERANS' satırının üstünde ay yazan bir satır olmalı.")
    header = [fold(v) if pd.notna(v) else "" for v in raw.iloc[hdr]]
    month_row = raw.iloc[hdr - 1].ffill()

    def find(pred):
        return next((i for i, h in enumerate(header) if pred(h)), None)

    c_ref = find(lambda h: h == "REFERANS")
    c_firma = find(lambda h: h == "FIRMA")
    c_backlog = find(lambda h: h == "BACKLOG")
    c_guncel = find(lambda h: "BACKLOG" in h and "GUNCEL" in h)
    c_stok = find(lambda h: h.startswith("STOK"))

    # --- ay blokları ve haftalar ---
    blocks: Dict[tuple, MonthBlock] = {}
    for i, h in enumerate(header):
        m = re.fullmatch(r"W\s*(\d{1,2})", h)
        if not m:
            continue
        hafta = int(m.group(1))
        if not 1 <= hafta <= 53:
            warnings.append(f"{h} geçersiz hafta numarası, atlandı.")
            continue
        parsed = parse_month_label(month_row.iloc[i])
        if parsed is None:
            raise ValueError(f"{h} haftasının üstünde ay adı okunamadı ('{month_row.iloc[i]}').")
        ay, ay_no, yil = parsed
        key = (ay, yil)
        if key not in blocks:
            blocks[key] = MonthBlock(label=str(month_row.iloc[i]).strip(), ay=ay, ay_no=ay_no, yil=yil)
        if hafta in blocks[key].weeks.values():
            warnings.append(f"{ay} içinde W{hafta} iki kez var; adetler toplandı.")
        blocks[key].weeks[i] = hafta
    if not blocks:
        raise ValueError("Hiç hafta sütunu (W34, W35 ...) bulunamadı.")
    block_list = list(blocks.values())

    # --- satırlar ---
    data = raw.iloc[hdr + 1:]
    num = lambda v: pd.to_numeric(v, errors="coerce")
    orders, backlog, refs = [], [], set()
    for _, row in data.iterrows():
        ref = normalize_code(row.iloc[c_ref])
        if not ref or "TOPLAM" in fold(ref):
            continue
        firma = str(row.iloc[c_firma]).strip() if c_firma is not None and pd.notna(row.iloc[c_firma]) else ""
        refs.add((ref, firma))
        for bi, b in enumerate(block_list):
            for col, hafta in b.weeks.items():
                q = num(row.iloc[col])
                if pd.notna(q) and q != 0:
                    if q < 0:
                        warnings.append(f"{ref} W{hafta}: negatif adet ({q:g}) atlandı.")
                        continue
                    orders.append((ref, firma, bi, hafta, float(q)))
        stok = num(row.iloc[c_stok]) if c_stok is not None else None
        stok = 0.0 if stok is None or pd.isna(stok) else float(stok)
        if c_backlog is not None and pd.notna(num(row.iloc[c_backlog])):
            bl = float(num(row.iloc[c_backlog]))
        elif c_guncel is not None and pd.notna(num(row.iloc[c_guncel])):
            bl = float(num(row.iloc[c_guncel])) + stok   # Güncel = Backlog - Stok
        else:
            bl = 0.0
        if bl or stok:
            backlog.append((ref, firma, bl, stok))

    # Formülü olup sonucu dosyada kayıtlı olmayan hücreler (dosya Excel'de hiç hesaplanıp kaydedilmemiş).
    # Bunlar okunursa 0 sayılır ve veri sessizce eksik yüklenir -> aktarma durdurulur.
    used_cols = sorted({c for b in block_list for c in b.weeks} | {c for c in (c_backlog, c_guncel, c_stok) if c is not None})
    missing = _uncached_formulas(path, sheet, raw, hdr, c_ref, used_cols)
    if missing:
        ornek = ", ".join(missing[:6])
        raise ValueError(
            f"Bu dosyada {len(missing)} hücrede formül var ama sonucu dosyaya kayıtlı değil (ör. {ornek}).\n"
            "Dosya Excel'de hiç açılıp kaydedilmemiş olabilir (başka bir programla oluşturulmuş dosyalarda olur).\n\n"
            "Çözüm: dosyayı Excel'de açın, Kaydet (Ctrl+S) deyin, kapatıp tekrar aktarın.\n"
            "Bu haliyle aktarılsaydı bu hücreler 0 sayılırdı; o yüzden hiçbir şey yazılmadı.")
    if not orders:
        warnings.append("Dosyada hiç sipariş adedi bulunamadı (tüm hafta hücreleri boş veya 0).")

    orders = pd.DataFrame(orders, columns=["urun_kodu", "firma", "blok", "hafta", "adet"])
    orders = orders.groupby(["urun_kodu", "firma", "blok", "hafta"], as_index=False)["adet"].sum()
    backlog = pd.DataFrame(backlog, columns=["urun_kodu", "firma", "backlog_adet", "stok_adet"])
    backlog = backlog.groupby(["urun_kodu", "firma"], as_index=False).sum()
    if c_backlog is None and c_guncel is None:
        warnings.append("Backlog sütunu bulunamadı; backlog aktarılmayacak.")
    return ParsedPlan(path, sheet, block_list, orders, backlog, len(refs), warnings)


def _uncached_formulas(path, sheet, raw, hdr, c_ref, cols) -> List[str]:
    """Okunan sütunlarda formülü olan ama hesaplanmış değeri olmayan hücrelerin adresleri (ör. 'H12')."""
    if not path.lower().endswith((".xlsx", ".xlsm")):
        return []   # .xls'te formül sonucu her zaman dosyaya yazılır; kontrol gerekmez
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    except Exception:
        return []
    try:
        ws = wb[sheet]
        out = []
        # pandas'taki satır r = Excel'de r+1. satır, sütun c = Excel'de c+1. sütun
        for r, row in enumerate(ws.iter_rows(min_row=hdr + 2, values_only=True), start=hdr + 1):
            if r >= len(raw) or not normalize_code(raw.iat[r, c_ref]):
                continue
            for c in cols:
                v = row[c] if c < len(row) else None
                is_formula = (isinstance(v, str) and v.startswith("=")) or type(v).__name__ == "ArrayFormula"
                if is_formula and pd.isna(raw.iat[r, c]):
                    out.append(f"{get_column_letter(c + 1)}{r + 1}")
        return out
    finally:
        wb.close()


def check_conflicts(db_path: str, plan: ParsedPlan, years: Dict[int, int]) -> List[str]:
    """Yüklemeden önce bilgi: aynı ay daha önce yüklenmiş mi, haftalar başka ayda var mı."""
    if not os.path.exists(db_path):
        return []
    create_empty_db(db_path)
    msgs = []
    con = connect(db_path)
    try:
        for bi, b in enumerate(plan.blocks):
            yil = years[bi]
            r = con.execute("SELECT id, tarih, dosya_adi FROM gecerli_aktarimlar WHERE yil=? AND ay=?",
                            (yil, b.ay)).fetchone()
            if r:
                msgs.append(f"{b.ay} {yil} daha önce yüklenmiş ({r[1]}, {r[2]}). Yeni yükleme geçerli olacak; "
                            "eskisi geçmişte kalır, istenirse geri alınabilir.")
            for hafta in sorted(set(b.weeks.values())):
                hy = hafta_yili(yil, b.ay_no, hafta)
                o = con.execute("SELECT DISTINCT ay || ' ' || yil FROM gecerli_siparisler "
                                "WHERE hafta_yili=? AND hafta=? AND NOT (ay=? AND yil=?)",
                                (hy, hafta, b.ay, yil)).fetchall()
                if o:
                    msgs.append(f"W{hafta} ({hy} yılının haftası) veritabanında {', '.join(x[0] for x in o)} "
                                "ayında da var; iki ayda birden sayılır.")
    finally:
        con.close()
    return msgs


def import_plan(db_path: str, plan: ParsedPlan, years: Dict[int, int]) -> List[int]:
    """Her ay bloğunu ayrı bir yükleme (işlem kaydı) olarak yazar. Tek işlem: hata olursa hiçbiri yazılmaz."""
    for bi, b in enumerate(plan.blocks):
        y = years.get(bi)
        if not isinstance(y, int) or not 2000 <= y <= 2100:
            raise ValueError(f"{b.ay} için geçerli bir yıl girilmedi.")
    create_empty_db(db_path)
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ids = []
    con = connect(db_path)
    try:
        with con:
            for bi, b in enumerate(plan.blocks):
                yil = years[bi]
                o = plan.orders[plan.orders["blok"] == bi]
                bl = plan.backlog if bi == 0 else plan.backlog.iloc[0:0]
                cur = con.execute(
                    "INSERT INTO aktarimlar(tarih, dosya_adi, sayfa, yil, ay, ay_no, haftalar, siparis_satir, backlog_satir) "
                    "VALUES (?,?,?,?,?,?,?,?,?)",
                    (now, os.path.basename(plan.path), plan.sheet, yil, b.ay, b.ay_no,
                     ",".join(str(h) for h in sorted(set(b.weeks.values()))), len(o), len(bl)))
                aid = cur.lastrowid
                con.executemany(
                    "INSERT INTO siparisler(aktarim_id, urun_kodu, firma, yil, hafta, hafta_yili, ay, adet) "
                    "VALUES (?,?,?,?,?,?,?,?)",
                    [(aid, r.urun_kodu, r.firma, yil, int(r.hafta), hafta_yili(yil, b.ay_no, int(r.hafta)), b.ay, r.adet)
                     for r in o.itertuples()])
                con.executemany(
                    "INSERT INTO backlog(aktarim_id, urun_kodu, firma, yil, ay, backlog_adet, stok_adet) "
                    "VALUES (?,?,?,?,?,?,?)",
                    [(aid, r.urun_kodu, r.firma, yil, b.ay, r.backlog_adet, r.stok_adet) for r in bl.itertuples()])
                ids.append(aid)
    finally:
        con.close()
    return ids


# ----------------------------------------------------------------------
# Operasyonlar (REFERANS PRES) – bir kerelik başlangıç verisi
# ----------------------------------------------------------------------
OP_SHEET_DEFAULT = "Uyum Operasyon Bazlı Saniyeler"
OP_HEADERS = {   # Excel başlığı (sadeleştirilmiş) -> veritabanı sütunu
    "MUSTERI": "musteri",
    "BITMIS URUN KODU": "urun_kodu",
    "OPERASYON STOK KODU": "operasyon_stok_kodu",
    "BASILDIGI PRES TONAJI": "pres_tonaji",
    "URETIM PRESI": "uretim_presi",
    "ALTERNATIF 1": "alternatif_1",
    "ALTERNATIF 2": "alternatif_2",
    "ALTERNATIF 3": "alternatif_3",
    "ALTERNATIF 4": "alternatif_4",
    "URETIM SANIYE": "uretim_saniye",
}
OP_REQUIRED = ["urun_kodu", "uretim_presi", "uretim_saniye"]


def parse_operations(path: str, sheet: str):
    """REFERANS PRES şablonundaki operasyon sayfasını okur. (DataFrame, uyarılar) döner."""
    from capacity_engine import normalize_machine
    from veritabani import OP_COLUMNS
    raw = read_sheet(path, sheet)
    hdr = None
    for r in range(min(20, len(raw))):
        if any(fold(v) == "BITMIS URUN KODU" for v in raw.iloc[r] if pd.notna(v)):
            hdr = r
            break
    if hdr is None:
        raise ValueError("Başlık satırı bulunamadı: ilk 20 satırda 'BİTMİŞ ÜRÜN KODU' yazan hücre yok.")
    cols = {}
    for i, v in enumerate(raw.iloc[hdr]):
        key = fold(v) if pd.notna(v) else ""
        if key in OP_HEADERS and OP_HEADERS[key] not in cols:
            cols[OP_HEADERS[key]] = i
    missing = [c for c in OP_REQUIRED if c not in cols]
    if missing:
        names = {v: k for k, v in OP_HEADERS.items()}
        raise ValueError("Şu sütunlar bulunamadı: " + ", ".join(names[c] for c in missing))

    miss_formula = _uncached_formulas(path, sheet, raw, hdr, cols["urun_kodu"], [cols["uretim_saniye"]])
    if miss_formula:
        raise ValueError(f"{len(miss_formula)} hücrede formül var ama sonucu kayıtlı değil (ör. "
                         f"{', '.join(miss_formula[:6])}). Dosyayı Excel'de açıp kaydedin, sonra tekrar deneyin.")

    rows, warnings, no_press, no_sec = [], [], [], []
    for _, row in raw.iloc[hdr + 1:].iterrows():
        kod = normalize_code(row.iloc[cols["urun_kodu"]])
        if not kod:
            continue
        rec = {}
        for c in OP_COLUMNS:
            v = row.iloc[cols[c]] if c in cols else None
            if c == "uretim_saniye":
                n = pd.to_numeric(v, errors="coerce")
                rec[c] = float(n) if pd.notna(n) and n > 0 else 0.0
            elif c == "urun_kodu":
                rec[c] = kod
            elif c == "operasyon_stok_kodu":
                rec[c] = normalize_code(v)
            elif c in ("uretim_presi", "alternatif_1", "alternatif_2", "alternatif_3", "alternatif_4"):
                rec[c] = normalize_machine(v)
            else:
                rec[c] = "" if v is None or (isinstance(v, float) and pd.isna(v)) else str(v).strip()
        if not rec["uretim_presi"]:
            no_press.append(kod)
        if rec["uretim_saniye"] <= 0:
            no_sec.append(kod)
        rows.append(rec)
    df = pd.DataFrame(rows, columns=OP_COLUMNS)
    if no_press:
        warnings.append(f"{len(no_press)} satırda Üretim presi boş (ör. {', '.join(no_press[:5])}).")
    if no_sec:
        warnings.append(f"{len(no_sec)} satırda ÜRETİM SANİYE boş veya 0 (ör. {', '.join(no_sec[:5])}).")
    dup = df.duplicated(["urun_kodu", "operasyon_stok_kodu"], keep=False) & (df["operasyon_stok_kodu"] != "")
    if dup.any():
        warnings.append(f"{int(dup.sum())} satırda aynı ürün + operasyon kodu birden fazla kez geçiyor.")
    return df, warnings
